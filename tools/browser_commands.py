import os
import re
import requests
from ddgs import DDGS
from docx import Document
from docx.shared import Cm


def search_internet(query: str) -> str:
    """Ищет актуальную информацию в интернете по любому запросу."""
    try:
        print(f"🌐 [Браузер]: Ищу в сети: '{query}'...")
        with DDGS() as ddgs:
            results = list(ddgs.text(query, max_results=5))
            if not results:
                return "Сэр, ничего не найдено."
            return "\n".join(
                [
                    f"Источник: {r['href']}\nТекст: {r['body']}\n"
                    for r in results
                ]
            )
    except Exception as e:
        return f"Ошибка поиска, сэр: {e}"


def create_word_document(
        topic: str, content: str, file_name: str = "report.docx"
) -> str:
    """Создает структурированный документ Word. Корректно разделяет текст и картинки."""
    try:
        if not file_name or not isinstance(file_name, str):
            file_name = "report.docx"

        file_name = file_name.replace('"', "").replace("'", "")

        desktop_path = os.path.join(os.path.expanduser("~"), "Desktop")
        full_file_path = os.path.join(desktop_path, file_name)
        print(f"📄 [Система]: Формирую продвинутый документ: {full_file_path}...")

        doc = Document()
        doc.add_heading(topic, level=1)

        lines = content.split("\n")

        with DDGS() as ddgs:
            for line in lines:
                line = line.strip()
                if not line:
                    continue

                # Ищем маркеры картинок в квадратных скобках [...]
                match = re.search(r"\[(.*?)\]", line)

                if match:
                    img_query = match.group(1).strip()

                    # Пропускаем, если это обычная ссылка или имя файла
                    if img_query.startswith("http") or img_query.endswith(
                            ".docx"
                    ):
                        doc.add_paragraph(line)
                        continue

                    # Вырезаем маркер картинки из строки, чтобы сохранить остальной текст
                    clean_text_line = line.replace(f"[{img_query}]", "").strip()

                    # Очищаем поисковый запрос от мусора, который мог оставить ИИ
                    search_query = img_query.replace("IMAGE:", "").replace("Картинка:", "").strip()
                    print(f"🖼️ [Система]: Скачиваю иллюстрацию: '{search_query}'")

                    try:
                        img_results = list(
                            ddgs.images(search_query, max_results=1)
                        )
                        if img_results:
                            img_url = img_results[0]["image"]
                            img_data = requests.get(img_url, timeout=10).content
                            temp_path = "temp_img.jpg"

                            with open(temp_path, "wb") as handler:
                                handler.write(img_data)

                            # Вставляем картинку (ширина 12 см)
                            doc.add_picture(temp_path, width=Cm(12.0))
                            os.remove(temp_path)

                            # Если в этой строке помимо картинки был текст, добавляем его абзацем ПОД картинкой
                            if clean_text_line:
                                doc.add_paragraph(clean_text_line)
                            continue

                    except Exception as img_err:
                        print(f"⚠️ Не удалось загрузить картинку '{search_query}': {img_err}")
                        # Если картинка упала, просто печатаем всю строку как текст
                        doc.add_paragraph(line)
                        continue

                # Обработка подзаголовков и обычного текста
                if line.startswith("###"):
                    doc.add_heading(line.replace("###", "").strip(), level=3)
                elif line.startswith("##"):
                    doc.add_heading(line.replace("##", "").strip(), level=2)
                else:
                    doc.add_paragraph(line)

        doc.add_paragraph(f"\nОтчет подготовлен вашим личным ассистентом Джарвисом.")
        doc.save(full_file_path)
        return f"Документ успешно создан на вашем Рабочем столе как '{file_name}', Серго!"

    except Exception as e:
        return f"Сэр, произошел сбой при сборке документа: {e}"





